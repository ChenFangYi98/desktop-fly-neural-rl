from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist" / "微信公众号文章 果蝇连接组具身强化学习.docx"


def set_font(run, name="Microsoft YaHei", size=11, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=130, start=140, bottom=130, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 11)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        set_font(r, size=16 if level == 1 else 13, bold=True)
    return p


def add_picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run(caption)
    set_font(r, size=9, color="666666")


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.72)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            set_font(r)
        if not p.runs:
            set_font(p.add_run(item))
        else:
            p.runs[0].text = item


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    style = styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(16)
title.paragraph_format.space_after = Pt(10)
set_font(title.add_run("我把果蝇连接组装进了三维世界"), size=25, bold=True)

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(20)
set_font(subtitle.add_run("一个可以训练技能并部署到具身机器人的开源实验室"), size=13, color="555555")

add_picture(doc, ROOT / "docs" / "images" / "ecosystem-brain.png",
            "生态空间 第一视角与实时大脑来自同一只果蝇")

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(14)
intro.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
set_font(intro.add_run("很多强化学习演示都有一个共同问题：智能体在地图里移动得很好，但我们不知道它的大脑在哪里。动作往往只是神经网络输出的几个数字，或者来自一张 Q 表。环境、策略、身体和可视化彼此分离，很难回答一个更有意思的问题：如果把真实连接组约束、神经活动、身体运动和三维生态放进同一个闭环里，会发生什么？"), size=12)

add_body(doc, "这正是我开源 DesktopFly Neural RL 的原因。")

add_heading(doc, "这不是给果蝇套一张 Q 表")
add_body(doc, "项目使用了两套公开连接组数据的派生子图。FlyWire v783 部分包含 668 个神经元和 18,968 条有符号连接，MaleCNS 部分包含 1,045 个神经元和 17,224 条运动相关连接。")
add_body(doc, "环境中的目标、食物、障碍和天敌首先被转换成感觉神经输入。活动沿有符号连接组传播，策略只能读取下降神经元活动，随后再进入 MaleCNS 和身体动力学。")
add_body(doc, "部署后的模型不能直接写入速度或航向。它必须走完环境感觉、FlyWire 神经活动、下降神经元、MaleCNS、身体和新环境感觉组成的闭环。这个限制让具身不再只是界面标签，而成为系统结构本身。")

add_picture(doc, ROOT / "assets" / "brain.png", "FlyWire 派生神经元位置与实时活动可视化")

add_heading(doc, "奖励如何改变大脑")
add_body(doc, "这次重构参考了 FlyDoom 的连接组强化学习方向，加入三因子学习。突触两端的神经活动形成局部资格迹，当前奖励和价值预测形成 TD 误差，TD 误差作为多巴胺调制信号。只有同时具备活动痕迹和多巴胺信号的连接才会发生变化。")
add_body(doc, "可以把它理解为：刚才活跃过的连接，会根据随后得到的奖励被增强或减弱。", bold_lead="可以把它理解为：")
add_body(doc, "当前版本支持三因子多巴胺学习、Connectome A2C 和 Connectome PPO。训练不会随意新增神经连接，也不会改变兴奋或抑制符号，而是在固定拓扑上学习有边界的群体突触增益。")

add_heading(doc, "地图和任务可以完全自定义")
add_picture(doc, ROOT / "docs" / "images" / "training-studio.png",
            "三维训练工作室 地形 素材 感觉 动作和奖惩都可编辑")
add_body(doc, "训练场景是一个真正的三维体素空间。用户可以选择草地、森林、峡谷等预设，也可以从空白地图开始。土块、岩石、水、食物、安全区、果蝇出生点、任务目标和天敌都能放到 X、Y、Z 任意合法坐标。")
add_body(doc, "青蛙、蜘蛛和螳螂拥有不同技能，并分为低、中、高三个捕食等级。任务规则也不是写死的。")
add_bullets(doc, [
    "观测可以包含位置、目标方向、碰撞、食物、天敌、能量、时间和脑活动",
    "动作空间可以组合六向移动和等待",
    "奖励可以来自接近目标、觅食、存活、远离天敌、高度或速度",
    "惩罚可以来自碰撞、能耗、被捕获或每一步成本",
    "回合可以在到达目标、吃到指定食物、存活足够时间或达到最大步数时结束",
])
add_body(doc, "因此，同一套基础连接组可以训练导航、觅食、逃生、高速机动或多任务技能。")

add_heading(doc, "模型有没有部署成功可以明确验证")
add_body(doc, "每次训练完成后，系统都会产生唯一模型编号，例如 DFM-20260921-153000-A1B2。模型包同时包含连接组指纹、改变了多少组连接、每个增益的数值、下降神经元读出、完整场景和奖励规则。")
add_body(doc, "在生态空间中，需要先选择模型文件，再选择目标果蝇。部署完成后，页面会显示模型编号、部署到第几号果蝇、训练增益映射到了多少条 LIF 突触、当前动作和累计决策次数。")
add_body(doc, "第一视角跟随哪只果蝇，大脑窗口就显示哪只果蝇。模型部署后，神经活动、动作和身体运动来自同一个运行时，不再是三个互不相关的动画。")

add_heading(doc, "从桌面生态到现实机器人")
add_body(doc, "导出的模型包包含 ROS 2 配置，可以由连接组群体运行时输出 Twist 指令，用于轮式机器人、六足平台或微型无人机的原型验证。")
add_body(doc, "现实机器人的电机和飞控不是果蝇肌肉。ROS 2 端执行同一连接组群体矩阵和下降神经读出，然后通过机器人动作适配器执行。真实部署仍需处理传感器噪声、定位漂移、动力学差异和时延。")
add_body(doc, "第一次真机测试必须使用独立硬件急停，限制速度，在架空或断开执行器的状态下检查方向，再进入封闭、低速、无人员环境。")

add_heading(doc, "如何理解真实果蝇大脑")
add_body(doc, "这个项目使用真实连接组派生数据，但它不是完整数字果蝇。真实的是神经元身份、位置、连接拓扑和突触计数；工程化的是感觉映射、LIF 参数、可塑性规则、跨标本接口、肌肉、飞行和机器人控制。")

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Cm(5.2)
table.columns[1].width = Cm(11.0)
headers = table.rows[0].cells
headers[0].text = "部分"
headers[1].text = "性质"
set_repeat_table_header(table.rows[0])
rows = [
    ("连接拓扑和突触计数", "公开连接组数据的派生子图"),
    ("神经动力学", "工程化 LIF 和连接组群体模型"),
    ("感觉编码", "三维环境到感觉神经群的工程适配器"),
    ("强化学习", "实际执行资格迹与 TD 多巴胺更新"),
    ("身体和机器人", "模型化身体动力学与硬件动作适配器"),
]
for cell in headers:
    set_cell_shading(cell, "203864")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for r in cell.paragraphs[0].runs:
        set_font(r, size=10, bold=True, color="FFFFFF")
for index, values in enumerate(rows):
    cells = table.add_row().cells
    for col, value in enumerate(values):
        cells[col].text = value
        cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[col])
        if index % 2:
            set_cell_shading(cells[col], "F2F6FA")
        for r in cells[col].paragraphs[0].runs:
            set_font(r, size=10)

add_body(doc, "为了让桌面电脑能够完成大量强化学习采样，训练阶段采用连接组群体可塑性：同一神经群之间的真实边共享一个学习增益。生态部署时，再把增益映射回 668 神经元 LIF 网络中的每一条对应边。")
add_body(doc, "更准确的表述是：这是一个连接组约束、神经活动可观察、奖励能够改变突触增益的具身强化学习实验平台。它离完整生物大脑还有很远，但已经比在策略网络旁边放一个大脑动画前进了一步。")

add_heading(doc, "为什么开源")
add_body(doc, "我希望它成为一个可以继续扩展的实验底座：增加更多感觉神经通道，尝试逐突触可塑性和更大连接组，接入更真实的空气动力学与六足动力学，并建立仿真到真实机器人的评估基准。")
add_body(doc, "项目源代码采用 MIT License。FlyWire 派生数据采用 CC BY-NC 4.0，只允许署名、非商业使用；MaleCNS 派生数据采用 CC BY 4.0。开源包中提供了数据来源、论文引用、模型边界、安装方法和自动化测试。")
add_body(doc, "如果你对连接组、强化学习、具身智能或机器人感兴趣，欢迎一起把它做得更扎实。")

add_heading(doc, "项目信息", level=2)
add_bullets(doc, [
    "项目名称  DesktopFly Neural RL",
    "关键词  果蝇连接组 强化学习 具身智能 三维生态 ROS 2",
    "适合读者  AI 神经科学 机器人 开源硬件和科学可视化开发者",
    "代码许可证  MIT",
    "数据许可证  FlyWire CC BY-NC 4.0  MaleCNS CC BY 4.0",
    "开源地址  上传 GitHub 后可在公众号原文链接中填写",
])

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "我把果蝇连接组装进了三维世界"
doc.core_properties.subject = "DesktopFly Neural RL 微信公众号文章"
doc.core_properties.author = "DesktopFly Neural RL"
doc.core_properties.keywords = "果蝇连接组 强化学习 具身智能 ROS 2"
doc.save(OUT)
print(OUT)

